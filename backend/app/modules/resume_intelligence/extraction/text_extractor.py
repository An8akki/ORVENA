"""
Resume Text Extractor

Extracts raw text from various resume formats (PDF, DOCX, TXT).
Handles different encodings and file formats gracefully.
"""
import io
import logging
from typing import Optional, Tuple
from pathlib import Path

# PDF parsing
try:
    import PyPDF2
    HAS_PDF_SUPPORT = True
except ImportError:
    HAS_PDF_SUPPORT = False

# DOCX parsing
try:
    import docx
    HAS_DOCX_SUPPORT = True
except ImportError:
    HAS_DOCX_SUPPORT = False

logger = logging.getLogger(__name__)


class ResumeTextExtractor:
    """
    Extracts text from resume files.
    
    Supports:
    - PDF files (via PyPDF2)
    - DOCX files (via python-docx)
    - TXT files (direct read)
    
    Falls back gracefully when libraries are unavailable.
    """
    
    def __init__(self):
        self.supported_formats = ['.txt']
        
        if HAS_PDF_SUPPORT:
            self.supported_formats.append('.pdf')
        
        if HAS_DOCX_SUPPORT:
            self.supported_formats.append('.docx')
    
    def extract_text(
        self,
        file_path: Optional[str] = None,
        file_content: Optional[bytes] = None,
        file_extension: Optional[str] = None
    ) -> Tuple[str, dict]:
        """
        Extract text from resume file.
        
        Args:
            file_path: Path to resume file
            file_content: File content as bytes (alternative to file_path)
            file_extension: File extension (required if using file_content)
            
        Returns:
            Tuple of (extracted_text, metadata)
            
        Raises:
            ValueError: If file format is unsupported or required args missing
        """
        # Determine file extension
        if file_path:
            ext = Path(file_path).suffix.lower()
        elif file_extension:
            ext = file_extension.lower()
            if not ext.startswith('.'):
                ext = f'.{ext}'
        else:
            raise ValueError("Either file_path or file_extension must be provided")
        
        # Check if format is supported
        if ext not in self.supported_formats:
            raise ValueError(
                f"Unsupported file format: {ext}. "
                f"Supported: {', '.join(self.supported_formats)}"
            )
        
        # Extract based on format
        if ext == '.pdf':
            return self._extract_from_pdf(file_path, file_content)
        elif ext == '.docx':
            return self._extract_from_docx(file_path, file_content)
        elif ext == '.txt':
            return self._extract_from_txt(file_path, file_content)
        else:
            raise ValueError(f"No extractor implemented for {ext}")
    
    def _extract_from_pdf(
        self,
        file_path: Optional[str],
        file_content: Optional[bytes]
    ) -> Tuple[str, dict]:
        """Extract text from PDF file."""
        if not HAS_PDF_SUPPORT:
            raise RuntimeError("PyPDF2 not installed. Install with: pip install PyPDF2")
        
        metadata = {'format': 'pdf', 'pages': 0}
        text_parts = []
        
        try:
            # Open PDF
            if file_path:
                with open(file_path, 'rb') as f:
                    pdf_reader = PyPDF2.PdfReader(f)
                    metadata['pages'] = len(pdf_reader.pages)
                    
                    for page in pdf_reader.pages:
                        text = page.extract_text()
                        if text:
                            text_parts.append(text)
            else:
                pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_content))
                metadata['pages'] = len(pdf_reader.pages)
                
                for page in pdf_reader.pages:
                    text = page.extract_text()
                    if text:
                        text_parts.append(text)
            
            full_text = '\n\n'.join(text_parts)
            metadata['char_count'] = len(full_text)
            metadata['success'] = True
            
            return full_text, metadata
            
        except Exception as e:
            logger.error(f"PDF extraction failed: {e}")
            metadata['success'] = False
            metadata['error'] = str(e)
            return "", metadata
    
    def _extract_from_docx(
        self,
        file_path: Optional[str],
        file_content: Optional[bytes]
    ) -> Tuple[str, dict]:
        """Extract text from DOCX file."""
        if not HAS_DOCX_SUPPORT:
            raise RuntimeError("python-docx not installed. Install with: pip install python-docx")
        
        metadata = {'format': 'docx', 'paragraphs': 0}
        text_parts = []
        
        try:
            # Open DOCX
            if file_path:
                doc = docx.Document(file_path)
            else:
                doc = docx.Document(io.BytesIO(file_content))
            
            # Extract paragraphs
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)
            
            # Extract from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text_parts.append(cell.text)
            
            full_text = '\n'.join(text_parts)
            metadata['paragraphs'] = len(doc.paragraphs)
            metadata['tables'] = len(doc.tables)
            metadata['char_count'] = len(full_text)
            metadata['success'] = True
            
            return full_text, metadata
            
        except Exception as e:
            logger.error(f"DOCX extraction failed: {e}")
            metadata['success'] = False
            metadata['error'] = str(e)
            return "", metadata
    
    def _extract_from_txt(
        self,
        file_path: Optional[str],
        file_content: Optional[bytes]
    ) -> Tuple[str, dict]:
        """Extract text from TXT file."""
        metadata = {'format': 'txt'}
        
        try:
            if file_path:
                with open(file_path, 'r', encoding='utf-8') as f:
                    text = f.read()
            else:
                # Try UTF-8 first, fall back to latin-1
                try:
                    text = file_content.decode('utf-8')
                except UnicodeDecodeError:
                    text = file_content.decode('latin-1')
            
            metadata['char_count'] = len(text)
            metadata['success'] = True
            
            return text, metadata
            
        except Exception as e:
            logger.error(f"TXT extraction failed: {e}")
            metadata['success'] = False
            metadata['error'] = str(e)
            return "", metadata
    
    def is_format_supported(self, file_extension: str) -> bool:
        """Check if a file format is supported."""
        ext = file_extension.lower()
        if not ext.startswith('.'):
            ext = f'.{ext}'
        return ext in self.supported_formats
